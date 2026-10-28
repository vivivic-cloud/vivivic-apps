#!/usr/bin/env python3
"""재단 제품목록 모으기 (사장님 10-28).

재단 명세서 xlsx 전부(날짜 시트 + 세부내역 시트)에서 제품명별 품목 정보를 중복 없이 모아
창고 `납품내역_제품목록` 에 담는다. 발주와 상관없는 단순 제품 정보다 — PO·수량·금액은 담지 않는다.
`제품목록.html` 은 이 창고를 읽기만 한다. 저절로 돌지 않는다 — 부를 때만 돈다.

부르는 법
  python3 납품내역/tools/제품목록모으기.py "<재단 xlsx 폴더>"          # 모아서 창고에 담음(통째로 갈아 끼움)
  python3 납품내역/tools/제품목록모으기.py "<재단 xlsx 폴더>" --안씀     # 모으기만(셈만 찍음)

  ⚠ 원본 xlsx 는 읽기만 한다. xlsx 여는 법·창고 문은 같은 폴더 `세부내역검사.py` 것을 그대로 쓴다.
  필요한 것: openpyxl
"""
import sys, os, re, glob, datetime, unicodedata, urllib.parse
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
import 세부내역검사 as 검사

창고 = 'artifacts/vivivic-4b7ef/public/data/납품내역_제품목록'
정보이름 = '_정보'
보일칸 = ['사이즈', '원장종류', '원장색상', '원장모델명', '원장출고처']
# PO 앞 글자 → 출고예정일 괄호 글자. 5월 파일에서 날짜 시트 「(ING)」 줄의 세부내역 PO 가 「본사…」 였다(16줄).
PO이름 = {'본사': 'ING', '조흔': '조혼'}


def 붙임(v):
    return re.sub(r'\s+', '', '' if v is None else str(v))


def 제품열쇠(v):
    # 날짜 시트에만 붙는 조 수 꼬리(「…(GIV) 36조」)는 뗀다.
    return re.sub(r'\d+조$', '', 붙임(v))


def 보임글(v):
    v = re.sub(r'\s+', ' ', '' if v is None else str(v)).strip()
    return '' if v in ('-', '－') else v


def 입고처(줄):
    """→ (과재고?, 괄호 글자 목록, 과재고 칸에 같이 적힌 PO 앞 글자 목록)"""
    예정, po = str(줄.get('출고예정일') or ''), str(줄.get('PO') or '')
    과 = '과재고' in 예정 or '과재고' in po
    괄호 = [x.strip() for x in re.findall(r'\(([^)]*)\)', 예정) if x.strip() and not re.search(r'\d', x)]
    앞 = [PO이름.get(x, x) for x in re.findall(r'([가-힣A-Za-z]+)\s*\d{4}-', po)]
    if 과:
        return True, [], 앞
    return False, 괄호 or 앞, []


def 모으기(폴더):
    파일들 = sorted(f for f in glob.glob(os.path.join(폴더, '*.xlsx')) if not 검사.파일이름(f).startswith('~$'))
    나온 = []          # (차례, 서류, 줄)
    시트수 = 0
    for f in 파일들:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for w in wb.worksheets:
            t = unicodedata.normalize('NFC', w.title).strip()
            if re.fullmatch(r'\d+일', t):
                줄들 = 검사.줄읽기(w.iter_rows(values_only=True), t)
            elif '세부내역' in t:
                줄들 = 검사.줄읽기(w.iter_rows(values_only=True))
            else:
                continue
            시트수 += 1
            서류 = f'{검사.파일이름(f)} · {t}'
            달 = 검사.달이름(f) or ''
            for 줄 in 줄들:
                if 제품열쇠(줄.get('제품명')) and 붙임(줄.get('품목')):
                    일 = int(re.match(r'\d+', str(줄.get('일') or '0')).group(0) or 0)
                    나온.append(((달, 일, len(나온)), 서류, 줄))

    제품 = {}           # 제품열쇠 → {이름들:Counter, 품목:{품목열쇠: 모음}}
    for 차례, 서류, 줄 in 나온:
        p = 제품.setdefault(제품열쇠(줄['제품명']), {'이름들': Counter(), '품목': {}})
        p['이름들'][re.sub(r'\s*\d+조$', '', 보임글(줄['제품명']))] += 1
        q = p['품목'].setdefault(붙임(줄['품목']), {'이름': Counter(), '칸': {k: {} for k in 보일칸},
                                                    '입고': {}, '과서류': set(), '같이': {}, '서류': set()})
        q['이름'][보임글(줄['품목'])] += 1
        q['서류'].add(서류)
        for k in 보일칸:
            v = 보임글(줄.get(k))
            if v:
                q['칸'][k][v] = max(q['칸'][k].get(v, 차례), 차례)
        과, 이름들, 같이 = 입고처(줄)
        if 과:
            q['과서류'].add(서류)
            for v in 같이:
                q['같이'][v] = max(q['같이'].get(v, 차례), 차례)
        for v in 이름들:
            q['입고'].setdefault(v, (차례, set()))
            q['입고'][v] = (max(q['입고'][v][0], 차례), q['입고'][v][1] | {서류})

    늦은차례 = lambda d: [v for v, _ in sorted(d.items(), key=lambda x: x[1], reverse=True)]
    목록, 셈 = [], Counter()
    for 열쇠 in sorted(제품, key=lambda k: 제품[k]['이름들'].most_common(1)[0][0]):
        p = 제품[열쇠]
        품목들 = []
        for q in p['품목'].values():
            칸 = {k: ' / '.join(늦은차례(q['칸'][k])) for k in 보일칸}
            입고 = {v: c for v, (c, _) in q['입고'].items()}
            값 = ' / '.join(늦은차례(입고))
            if q['과서류']:
                셈['과재고줄'] += 1
                if 값:          # 같은 제품·품목이 과재고가 아니었던 다른 줄(앞뒤 서류)의 괄호 글자
                    셈['메움'] += 1
                elif not 값 and q['같이']:
                    값 = ' / '.join(늦은차례(q['같이']))
                    셈['같은칸메움'] += 1
                elif not 값:
                    셈['못메움'] += 1
            if not 값 and not q['과서류']:
                셈['괄호없음'] += 1
            품목들.append({'품목': q['이름'].most_common(1)[0][0], **칸, '출고예정일': 값})
        목록.append({'제품명': p['이름들'].most_common(1)[0][0], '품목들': 품목들})
    return 목록, {'파일수': len(파일들), '시트수': 시트수, '읽은줄': len(나온), '제품수': len(목록),
                 '줄수': sum(len(x['품목들']) for x in 목록), **셈}


def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    폴더, 안씀 = sys.argv[1], '--안씀' in sys.argv
    목록, 셈 = 모으기(폴더)
    for k, v in 셈.items():
        print(f'{k}: {v}')
    if 안씀:
        return
    # ⚠ 창고문 의 쓰기·읽기 자리는 세부내역검사.py 의 전역 `창고`(= 납품내역_세부내역)다.
    #    바꿔 끼우지 않으면 제품 문서가 세부내역 창고에 들어간다(10-28 한 번 그렇게 됐다).
    검사.창고 = 창고
    문 = 검사.창고문()
    문서들 = {f'제품_{i:04d}': {'순번': i, **x} for i, x in enumerate(목록, 1)}
    문.여럿쓰기(문서들)
    # 지난번보다 제품이 줄었으면 남은 문서를 지운다(이 창고 안에서만).
    있던 = 문._부름(f'{문.뿌리}/{urllib.parse.quote(창고)}?pageSize=1000') or {}
    for d in 있던.get('documents', []):
        n = d['name'].rsplit('/', 1)[1]
        if n.startswith('제품_') and n not in 문서들:
            문._부름(f'{문.뿌리}/{문._자리(n)}', method='DELETE')
    때 = datetime.datetime.now().astimezone().isoformat(timespec='seconds')
    문.여럿쓰기({정보이름: {'모은때': 때, **셈}})
    print(f'→ 창고에 제품 {len(목록)}개 문서 + {정보이름}')


if __name__ == '__main__':
    main()
