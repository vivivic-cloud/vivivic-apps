#!/usr/bin/env python3
"""재단 세부내역 일치 검사 (사장님 10-22).

재단 xlsx 한 달 파일 안의
  날짜 시트(1일·2일…)의 재단 줄을 다 합친 것  ==  「N월 세부내역」 시트의 줄
인지 본다. 일치하는 달만 창고에 세부내역 줄 + 「검사 끝난 달」 표시를 남긴다.
검사 끝난 달은 다음에 부르면 건너뛴다. 저절로 돌지 않는다 — 부를 때만 돈다.

부르는 법
  python3 납품내역/tools/세부내역검사.py "<재단 xlsx 폴더>"          # 검사 + 일치한 달 기록
  python3 납품내역/tools/세부내역검사.py "<재단 xlsx 폴더>" --안씀     # 검사만(창고에 안 씀)

  ⚠ 원본 xlsx 는 읽기만 한다. 드라이브 경로·xlsx 내용은 저장소에 두지 않는다.
  창고 열쇠는 저장소에 새로 박지 않고, 같은 폴더 `../재단.html` 의 웹 설정(이미 공개)에서 읽는다.
  필요한 것: openpyxl (pip3 install openpyxl)
"""
import sys, os, re, glob, json, datetime, unicodedata, urllib.request, urllib.parse
from collections import Counter

import openpyxl

# 잣대 — 이 칸이 모두 같으면 같은 줄로 본다(줄 순서는 보지 않고 묶음 수로 맞춘다).
#   글은 띄어쓰기·줄바꿈을 다 빼고 견준다(「W15조」 = 「W 15조」).
잣대칸 = ['일', '제품명', '원장종류', '원장색상', '원장모델명', '원장출고처',
         '품목', '사이즈', '재단수량', '원장수량', '재단금액']
# 참고칸 — 세부내역을 만든 뒤 고쳐 적는 메모 칸이라(예: 「빠른일정」 → 「4/17 (ING)」) 일치에 넣지 않는다.
#   다른 곳은 「참고」 로 세어 보이기만 한다.
참고칸 = ['원장입고일', '출고예정일', '비고']
# 칸 자리는 시트마다 머리줄 글로 잡는다 — 5월부터 세부내역에만 「PO NO.」 칸이 끼어 한 칸씩 밀린다.
머리글 = {'제품명': '제품명', 'PONO.': 'PO', '원장입고일': '원장입고일', '품목': '품목', '사이즈': '사이즈',
         '재단수량': '재단수량', '원장수량': '원장수량', '재단금액': '재단금액', '비고': '비고'}
원장글 = {'종류': '원장종류', '색상': '원장색상', '모델명': '원장모델명', '출고처': '원장출고처'}
# 이어 붙이는 칸 — 첫 줄에만 적히고 아래 품목 줄은 빈칸이다.
제품칸 = ['제품명', 'PO', '출고예정일', '비고']                      # 새 제품명이 나오면 새로 잡는다
원장칸 = ['원장종류', '원장색상', '원장모델명', '원장출고처', '원장입고일']  # 한 제품 안에서도 바뀐다(예: 뒷판만 MDF-3T)
담는칸 = ['제품명', 'PO'] + 원장칸 + ['품목', '사이즈', '재단수량', '원장수량', '재단금액', '출고예정일', '비고']

# 세부내역 줄은 납품내역 창고 옆 칸에 둔다 — 재단.html 이 납품내역 창고를 달마다 통째로 받아서,
# 같은 창고에 넣으면 그 쪽이 그만큼 느려진다. 줄 꼴(달·갈·일·서류 …)은 납품내역 창고와 같게.
창고 = 'artifacts/vivivic-4b7ef/public/data/납품내역_세부내역'
표시이름 = '_검사'                 # 「검사 끝난 달」 표시 문서(달마다 칸 하나)


def 다듬(v):
    """엑셀 칸 값을 견주기 좋게 — 글은 줄바꿈·띄어쓰기를 한 칸으로, 숫자는 정수면 정수."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = re.sub(r'\s+', ' ', v).strip()
        return v or None
    if isinstance(v, (datetime.date, datetime.datetime)):
        return f'{v.month}/{v.day}'
    return v


def 붙임(v):
    return '' if v is None else re.sub(r'\s+', '', str(v))


def 칸잡기(머리1, 머리2):
    """머리줄 둘 → {칸이름: 자리}"""
    자리 = {}
    for i, v in enumerate(머리1):
        k = 붙임(v)
        if k in 머리글:
            자리[머리글[k]] = i
        elif k.startswith('출고예정일'):
            자리['출고예정일'] = i
    for i, v in enumerate(머리2):
        if 붙임(v) in 원장글:
            자리[원장글[붙임(v)]] = i
    return 자리


def 줄읽기(rows, 날짜시트일=None):
    """시트 줄들 → 재단 줄 목록. 날짜시트일이 있으면 날짜 시트(일은 시트 이름), 없으면 세부내역(첫 칸이 재단일)."""
    out, 제품, 원장, 일, 자리, 앞줄 = [], {}, {}, 날짜시트일, None, None
    for r in rows:
        r = list(r) + [None] * 20
        a글 = 붙임(r[0])
        # 재단 표가 끝나는 자리: 날짜 시트는 「합계」, 세부내역은 아래 「출고처…총소요량」 표.
        if a글.startswith('합계') or a글 in ('출고처', '총소요량'):
            break
        if a글 in ('순번', '재단일'):        # 머리 첫 줄(세부내역은 쪽마다 되풀이) — 다음 줄과 함께 칸을 잡는다
            앞줄 = r
            continue
        if 앞줄 is not None:                  # 머리 둘째 줄(종류·색상·모델명·출고처)
            자리, 앞줄 = 칸잡기(앞줄, r), None
            continue
        if 자리 is None:
            continue
        if 날짜시트일 is None and re.fullmatch(r'\d+일', a글):
            일 = a글
        값 = {k: 다듬(r[i]) for k, i in 자리.items()}
        if 값.get('품목') is None:
            continue
        if 값.get('제품명') is not None:      # 새 제품
            제품 = {k: 값.get(k) for k in 제품칸}
            원장 = {}
        if any(값.get(k) is not None for k in 원장칸):
            원장 = {k: 값.get(k) if 값.get(k) is not None else 원장.get(k) for k in 원장칸}
        줄 = {'일': 일}
        for k in 담는칸:
            줄[k] = 제품.get(k) if k in 제품칸 else 원장.get(k) if k in 원장칸 else 값.get(k)
        out.append(줄)
    return out


def 달파일읽기(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    세부 = [w for w in wb.worksheets if '세부내역' in unicodedata.normalize('NFC', w.title)]
    날짜 = [w for w in wb.worksheets if re.fullmatch(r'\d+일', unicodedata.normalize('NFC', w.title).strip())]
    세부줄 = 줄읽기(세부[0].iter_rows(values_only=True)) if 세부 else None
    날짜줄 = []
    for w in 날짜:
        날짜줄 += 줄읽기(w.iter_rows(values_only=True), unicodedata.normalize('NFC', w.title).strip())
    return {'세부시트': 세부[0].title if 세부 else None, '날짜시트수': len(날짜),
            '세부줄': 세부줄, '날짜줄': 날짜줄}


def 견줄값(v, 칸=None):
    # 같은 수가 어떤 시트에는 글('28800'), 어떤 시트에는 수(28800)로 적혀 있다 — 수로 맞춘다.
    # 「-」 는 빈칸과 같은 뜻으로 본다(재단금액이 한쪽은 「-」, 한쪽은 빈칸인 줄이 있다).
    # 제품명 끝의 조 수(「…(GIV) 36조」)는 날짜 시트에만 붙는 꼬리라 떼고 견준다.
    if isinstance(v, str):
        v = re.sub(r'\s+', '', v)
        if v == '-':
            return None
        if 칸 == '제품명':
            v = re.sub(r'\d+조$', '', v)
        if re.fullmatch(r'-?\d+(\.\d+)?', v):
            f = float(v)
            return int(f) if f.is_integer() else f
    return v


def 열쇠(줄, 칸=None):
    return tuple(json.dumps(견줄값(줄.get(k), k), ensure_ascii=False) for k in (칸 or 잣대칸))


def 견주기(d):
    """→ (일치?, 까닭 글 목록)"""
    if d['세부줄'] is None:
        return False, ['세부내역 시트가 없다']
    if d['날짜시트수'] == 0:
        return False, ['날짜 시트가 없다 — 견줄 것이 없어 검사할 수 없다']
    a, b = Counter(map(열쇠, d['날짜줄'])), Counter(map(열쇠, d['세부줄']))
    전칸 = 잣대칸 + 참고칸
    참고 = sum((Counter(열쇠(x, 전칸) for x in d['날짜줄']) - Counter(열쇠(x, 전칸) for x in d['세부줄'])).values())
    참고글 = [f'참고: 메모 칸({"·".join(참고칸)})만 다른 줄 {참고}개 — 일치에 넣지 않음'] if 참고 else []
    if a == b:
        return True, 참고글
    글 = [f'줄 수: 날짜 시트 합 {len(d["날짜줄"])} · 세부내역 {len(d["세부줄"])}'
         f' · 날짜 시트에만 {sum((a - b).values())} · 세부내역에만 {sum((b - a).values())}']
    for 이름, 차 in (('날짜 시트에만', a - b), ('세부내역에만', b - a)):
        for k, n in sorted(차.items()):
            v = dict(zip(잣대칸, (json.loads(x) for x in k)))
            글.append(f'  {이름} ×{n}: ' + ' · '.join(f'{c}={v[c]}' for c in 잣대칸 if v[c] is not None))
    return False, 글


def 파일이름(path):
    # 맥 드라이브는 한글 이름을 풀어쓴 꼴(NFD)로 준다 — 모아쓴 꼴(NFC)로 맞춘다(창고 `서류` 칸과 같게).
    return unicodedata.normalize('NFC', os.path.basename(path))


def 달이름(path):
    m = re.search(r'(\d+)년\s*(\d+)월', 파일이름(path))
    return f'20{int(m.group(1)):02d}-{int(m.group(2)):02d}' if m else None


# ───────── 창고(Firestore REST) ─────────
class 창고문:
    def __init__(self):
        html = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '재단.html'), encoding='utf-8').read()
        웹설정 = re.search(r'apiKey:\s*"([^"]+)"', html).group(1)
        r = self._부름(f'https://identitytoolkit.googleapis.com/v1/accounts:signUp?key={웹설정}',
                      {'returnSecureToken': True}, 토큰=None)
        self.토큰 = r['idToken']
        self.뿌리 = 'https://firestore.googleapis.com/v1/projects/vivivic-4b7ef/databases/(default)/documents'

    def _부름(self, url, body=None, 토큰='', method=None):
        h = {'content-type': 'application/json'}
        if 토큰 is not None:
            h['Authorization'] = 'Bearer ' + self.토큰
        req = urllib.request.Request(url, data=None if body is None else json.dumps(body).encode(),
                                     headers=h, method=method or ('POST' if body is not None else 'GET'))
        try:
            with urllib.request.urlopen(req, timeout=60) as res:
                return json.loads(res.read() or b'{}')
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return None
            raise RuntimeError(f'{e.code} {e.read().decode()[:300]}')

    def _자리(self, 이름):
        return urllib.parse.quote(f'{창고}/{이름}')

    def 읽기(self, 이름):
        r = self._부름(f'{self.뿌리}/{self._자리(이름)}')
        return None if r is None else {k: 값풀기(v) for k, v in r.get('fields', {}).items()}

    def 칸더하기(self, 이름, 칸들):
        """문서에 칸만 덧붙인다(다른 칸은 그대로 — updateMask)."""
        q = '&'.join('updateMask.fieldPaths=' + urllib.parse.quote(f'`{k}`') for k in 칸들)
        self._부름(f'{self.뿌리}/{self._자리(이름)}?{q}', {'fields': {k: 값싸기(v) for k, v in 칸들.items()}},
                   method='PATCH')

    def 여럿쓰기(self, 문서들):
        """{이름: 칸들} 을 400개씩 묶어 통째로 쓴다(같은 이름이면 덮어씀 — 다시 돌려도 줄이 늘지 않는다)."""
        이름들 = list(문서들)
        for i in range(0, len(이름들), 400):
            쓰기 = [{'update': {'name': f'projects/vivivic-4b7ef/databases/(default)/documents/{창고}/{n}',
                                'fields': {k: 값싸기(v) for k, v in 문서들[n].items()}}} for n in 이름들[i:i + 400]]
            self._부름(f'{self.뿌리}:commit', {'writes': 쓰기})


def 값싸기(v):
    if v is None: return {'nullValue': None}
    if isinstance(v, bool): return {'booleanValue': v}
    if isinstance(v, int): return {'integerValue': str(v)}
    if isinstance(v, float): return {'doubleValue': v}
    if isinstance(v, dict): return {'mapValue': {'fields': {k: 값싸기(x) for k, x in v.items()}}}
    if isinstance(v, list): return {'arrayValue': {'values': [값싸기(x) for x in v]}}
    return {'stringValue': str(v)}


def 값풀기(v):
    (t, x), = v.items()
    if t == 'integerValue': return int(x)
    if t == 'mapValue': return {k: 값풀기(y) for k, y in x.get('fields', {}).items()}
    if t == 'arrayValue': return [값풀기(y) for y in x.get('values', [])]
    if t == 'nullValue': return None
    return x


def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    폴더, 안씀 = sys.argv[1], '--안씀' in sys.argv
    파일들 = sorted(f for f in glob.glob(os.path.join(폴더, '*.xlsx')) if not 파일이름(f).startswith('~$'))
    문 = None if 안씀 else 창고문()
    끝난 = (문.읽기(표시이름) or {}) if 문 else {}
    for f in 파일들:
        달 = 달이름(f)
        if not 달:
            continue
        if 끝난.get(달, {}).get('일치') is True:
            print(f'{달}  건너뜀 — {끝난[달].get("검사한때")} 에 검사 끝남({끝난[달].get("줄수")}줄)')
            continue
        d = 달파일읽기(f)
        맞음, 까닭 = 견주기(d)
        n세부 = len(d['세부줄'] or [])
        print(f'{달}  {"일치" if 맞음 else "불일치"} — 세부내역 {n세부}줄 · 날짜 시트 {d["날짜시트수"]}장 {len(d["날짜줄"])}줄')
        for g in 까닭:
            print('   ', g)
        if not 맞음 or 안씀:
            continue
        서류 = 파일이름(f)
        문서들 = {}
        for i, 줄 in enumerate(d['세부줄'], 1):
            문서들[f'세부내역_{달}_{i:04d}'] = {'달': 달, '갈': '세부내역', '서류': 서류,
                                             '시트': d['세부시트'], '번호': i, **줄}
        문.여럿쓰기(문서들)
        때 = datetime.datetime.now().astimezone().isoformat(timespec='seconds')
        문.칸더하기(표시이름, {달: {'일치': True, '검사한때': 때, '줄수': n세부,
                                 '날짜시트수': d['날짜시트수'], '서류': 서류, '잣대': 잣대칸}})
        print(f'    → 창고에 세부내역 {n세부}줄 + 검사 끝남 표시를 남겼다')


if __name__ == '__main__':
    main()
